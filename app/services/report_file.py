from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.constants import ReportType
from app.services.report_generation import ReportGeneration


class ReportFile:
    def __init__(self):

        self.sheets = {
            ReportType.REPORT_OVER_POSITION.name: {
                "tax_id": "РНОКПП військовослужбовця",
                "date_raport": "Дата складання рапорту",
            },
            ReportType.REPORT_ACCEPTED_POSITION.name: {
                "position_code": "Код посади",
                "tax_id": "РНОКПП військовослужбовця",
                "order_name": "Назва наказу",
                "oder_number": "Номер наказу",
                "oder_date": "Дата наказу",
                "date_raport": "Дата складання рапорту",
            },
        }

    def create_excel_template(self, filename: str):
        wb = Workbook()

        default_sheet = wb.active
        wb.remove(default_sheet)

        for sheet_name, fields in self.sheets.items():

            ws = wb.create_sheet(title=sheet_name)

            for col, (field_name, hint) in enumerate(fields.items(), start=1):

                column_letter = get_column_letter(col)

                cell = ws.cell(row=1, column=col, value=field_name)

                cell.font = Font(bold=True)

                cell.alignment = Alignment(horizontal="center", vertical="center")

                ws.column_dimensions[column_letter].width = max(len(field_name) + 3, 15)

                dv = DataValidation(
                    type="textLength",
                    operator="lessThanOrEqual",
                    formula1="1000",
                    allow_blank=True,
                )

                dv.promptTitle = field_name
                dv.prompt = hint

                dv.errorTitle = "Помилка"
                dv.error = f"Перевірте значення поля: {field_name}"

                dv.showInputMessage = True
                dv.showErrorMessage = True

                ws.add_data_validation(dv)

                dv.add(f"{column_letter}2:" f"{column_letter}1000")

            ws.freeze_panes = "A2"

            last_column = get_column_letter(len(fields))

            ws.auto_filter.ref = f"A1:{last_column}1"

        wb.save(filename)

    def load_excel(self, filename: str) -> dict:
        wb = load_workbook(filename=filename, data_only=True)

        self._validate_sheets(wb)

        result = {}

        for sheet_name, fields in self.sheets.items():

            ws = wb[sheet_name]

            headers = [cell.value for cell in ws[1] if cell.value is not None]

            self._validate_headers(
                sheet_name=sheet_name, headers=headers, expected_fields=fields
            )

            rows = []

            for row in ws.iter_rows(min_row=2, values_only=True):
                if all(value is None for value in row):
                    continue

                row_data = {}

                for index, field_name in enumerate(headers):
                    row_data[field_name] = row[index]

                rows.append(row_data)

            result[sheet_name] = rows

        return result

    def _validate_sheets(self, wb):
        expected_sheets = set(self.sheets.keys())
        actual_sheets = set(wb.sheetnames)

        missing = expected_sheets - actual_sheets
        extra = actual_sheets - expected_sheets

        if missing:
            raise ValueError(f"Відсутні листи: {', '.join(map(str, missing))}")

        if extra:
            raise ValueError(f"Невідомі листи: {', '.join(map(str, extra))}")

    def _validate_headers(self, sheet_name: str, headers: list, expected_fields: dict):
        expected = set(expected_fields.keys())
        actual = set(headers)

        missing = expected - actual
        extra = actual - expected

        if missing:
            raise ValueError(
                f"Лист '{sheet_name}': " f"відсутні поля: {', '.join(missing)}"
            )

        if extra:
            raise ValueError(
                f"Лист '{sheet_name}': " f"невідомі поля: {', '.join(extra)}"
            )

        if len(headers) != len(set(headers)):
            raise ValueError(f"Лист '{sheet_name}': " f"є дубльовані заголовки")
